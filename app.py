# -*- coding: utf-8 -*-
"""
RAG 网页版（Streamlit + bge + FAISS + DeepSeek）
上传文档 → 自动建立语义索引 → 提问 → 显示答案 + 依据的原文片段

首次运行前安装（PyCharm 底部 Terminal）：
    pip install streamlit sentence-transformers faiss-cpu openai pdfplumber python-docx -i https://pypi.tuna.tsinghua.edu.cn/simple

运行（在 learn_ai 目录下）：
    streamlit run app.py
"""

import os
os.environ["HF_ENDPOINT"] = "https://hf-mirror.com"   # 国内下载 bge 模型加速

import io
import re
import faiss
import pdfplumber
import docx
import streamlit as st
from sentence_transformers import SentenceTransformer
from openai import OpenAI

# ================= 配置 =================
EMBED_MODEL_NAME = "BAAI/bge-small-zh-v1.5"
CHUNK_SIZE = 400
CHUNK_OVERLAP = 50

# ================= 页面 =================
st.set_page_config(page_title="RAG 知识问答", page_icon="📚", layout="wide")
st.title("📚 RAG 知识问答系统")
st.caption("上传文档 → 建立语义索引 → 提问，答案会标注它依据的原文片段")

with st.sidebar:
    st.header("⚙️ 设置")
    # 优先从 Streamlit Cloud 的 Secrets 读取 key；本地没配置则显示输入框
    secret_key = ""
    try:
        secret_key = st.secrets.get("DEEPSEEK_API_KEY", "")
    except Exception:
        pass
    if secret_key:
        api_key = secret_key
        st.success("✅ 已使用云端配置的 API Key")
    else:
        api_key = st.text_input("DeepSeek API Key", type="password",
                                placeholder="sk-...（本地运行填这里）")
    top_k = st.slider("检索片段数 TOP_K", 1, 8, 4)
    st.markdown("---")
    st.caption("第一次运行会自动下载 bge 模型，耐心等")

# ================= 加载向量模型（缓存，避免每次交互都重新加载） =================
@st.cache_resource
def get_embedder():
    return SentenceTransformer(EMBED_MODEL_NAME)

# ================= 读取上传的文件 =================
def read_uploaded_file(file):
    """根据文件类型读取文字，返回字符串"""
    name = file.name.lower()
    data = file.getvalue()
    if name.endswith(".txt"):
        return data.decode("utf-8")
    if name.endswith(".pdf"):
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    if name.endswith(".docx"):
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts)
    return ""

# ================= 切分 =================
def chunk_text(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    sentences = re.split(r"(?<=[。！？!?])", text)
    sentences = [s.strip() for s in sentences if s.strip()]
    chunks, cur = [], ""
    for s in sentences:
        if cur and len(cur) + len(s) > chunk_size:
            chunks.append(cur)
            cur = cur[-overlap:] + s
        else:
            cur += s
    if cur:
        chunks.append(cur)
    return chunks

# ================= 建索引 =================
def build_index(docs):
    """docs: [(文件名, 全文), ...] → 返回 (index, chunks)"""
    embedder = get_embedder()
    chunks = []   # 每个元素 {"file": 文件名, "text": 片段}
    for fname, text in docs:
        for c in chunk_text(text):
            if c.strip():
                chunks.append({"file": fname, "text": c})
    vectors = embedder.encode([c["text"] for c in chunks], normalize_embeddings=True)
    dim = vectors.shape[1]
    index = faiss.IndexFlatIP(dim)
    index.add(vectors)
    return index, chunks

def retrieve(index, chunks, query, top_k):
    embedder = get_embedder()
    q_vec = embedder.encode([query], normalize_embeddings=True)
    scores, ids = index.search(q_vec, top_k)
    results = []
    for score, i in zip(scores[0], ids[0]):
        if i >= 0:
            results.append((score, chunks[i]))
    return results

def generate(client, query, results):
    context = "\n\n".join(f"[资料{idx+1}] {r[1]['text']}" for idx, r in enumerate(results))
    prompt = f"""请只根据下面提供的资料回答问题，资料里没有的内容就明确说不知道，不要编造。

资料：
{context}

问题：{query}

回答："""
    resp = client.chat.completions.create(
        model="deepseek-chat",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.3,
    )
    return resp.choices[0].message.content

# ================= 主界面 =================
uploaded = st.file_uploader("上传文档（txt / pdf / docx，可多选）",
                            type=["txt", "pdf", "docx"],
                            accept_multiple_files=True)

if st.button("📥 建立索引", type="primary", disabled=not uploaded):
    with st.spinner("正在读取文档、建立语义索引……"):
        docs = []
        for f in uploaded:
            text = read_uploaded_file(f)
            if text.strip():
                docs.append((f.name, text))
            else:
                st.warning(f"{f.name} 没提取到文字（可能是扫描版 PDF），已跳过")
        if not docs:
            st.error("没有可用的文字，请换文字版文档（能选中复制文字的才算）")
        else:
            index, chunks = build_index(docs)
            st.session_state["index"] = index
            st.session_state["chunks"] = chunks
            st.success(f"✅ 已入库 {len(chunks)} 个文本片段，现在可以提问了")

# 提问区
if "index" in st.session_state:
    st.markdown("---")
    query = st.text_input("请输入问题：")
    if st.button("🔍 提问", disabled=not (query and api_key)):
        if not api_key:
            st.error("请先在左侧填 DeepSeek API Key")
        else:
            client = OpenAI(api_key=api_key, base_url="https://api.deepseek.com")
            results = retrieve(st.session_state["index"], st.session_state["chunks"], query, top_k)
            with st.spinner("生成答案中……"):
                answer = generate(client, query, results)

            st.markdown("### 回答")
            st.write(answer)

            st.markdown("### 依据的原文片段")
            for idx, (score, c) in enumerate(results):
                with st.expander(f"片段 {idx+1}（相似度 {score:.3f} · 来自《{c['file']}》）"):
                    st.write(c["text"])
