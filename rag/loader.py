# -*- coding: utf-8 -*-
"""文档加载：读取 txt / pdf / docx，返回 (文件名, 文本) 列表。

pdfplumber / python-docx 属于重依赖，只在真正读对应格式时才 import，
保证 import rag 本身不触发安装这些库。
"""
import glob
import logging
import os

logger = logging.getLogger("rag.loader")


def read_txt(path):
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


def read_pdf(path):
    try:
        import pdfplumber
    except ImportError:
        raise ImportError("缺少 pdfplumber，请先运行：pip install pdfplumber")
    parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                parts.append(t)
    return "\n".join(parts)


def read_docx(path):
    try:
        import docx
    except ImportError:
        raise ImportError("缺少 python-docx，请先运行：pip install python-docx")
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:            # 顺便把表格里的文字也读出来
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def read_file(path):
    """根据扩展名读取单个文件，返回文本。"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".docx":
        return read_docx(path)
    return read_txt(path)


def read_file_bytes(name, data):
    """从内存字节读取文档（网页 / API 上传用），返回文本。"""
    name = name.lower()
    if name.endswith(".txt"):
        return data.decode("utf-8")
    if name.endswith(".pdf"):
        import io
        try:
            import pdfplumber
        except ImportError:
            raise ImportError("缺少 pdfplumber，请先运行：pip install pdfplumber")
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    if name.endswith(".docx"):
        import io
        try:
            import docx
        except ImportError:
            raise ImportError("缺少 python-docx，请先运行：pip install python-docx")
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts)
    return ""


def load_docs_folder(folder):
    """读取文件夹里所有 .txt / .pdf / .docx，返回 [(文件名, 文本), ...]。"""
    paths = sorted(
        glob.glob(os.path.join(folder, "*.txt"))
        + glob.glob(os.path.join(folder, "*.pdf"))
        + glob.glob(os.path.join(folder, "*.docx"))
    )
    docs = []
    for path in paths:
        name = os.path.basename(path)
        try:
            text = read_file(path).strip()
            if text:
                docs.append((name, text))
            else:
                logger.warning("跳过（没提取到文字，可能是扫描版 PDF）：%s", name)
        except Exception as e:
            logger.warning("读取失败 %s：%s", name, e)
    logger.info("已读取 %d 个文档：%s", len(docs), [n for n, _ in docs])
    return docs
