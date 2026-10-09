# -*- coding: utf-8 -*-
"""
扫描版 PDF → Word（OCR 文字识别）
原理：用 pymupdf 把每一页渲染成图片，再用 RapidOCR 识别图片里的文字，最后写进 Word。

使用前安装依赖（一次性，在 PyCharm 底部 Terminal 里粘贴）：
    pip install pymupdf rapidocr_onnxruntime python-docx -i https://pypi.tuna.tsinghua.edu.cn/simple

用法：
    1. 把下面 PDF_PATH 改成你的 PDF 路径（默认已填好）
    2. 第一次运行会自动下载 OCR 模型（几十 MB），耐心等
    3. 强烈建议：先只转 2~3 页看效果（把 END_PAGE 设小），别一上来就整本转
"""

import fitz  # pymupdf：把 PDF 页面渲染成图片
from rapidocr_onnxruntime import RapidOCR
from docx import Document

# ================= 配置 =================
PDF_PATH = r"C:\Users\Administrator\Desktop\learn_ai\docs\何曼君《高分子物理》（第3版）.pdf"
OUT_PATH = r"C:\Users\Administrator\Desktop\高分子物理_识别结果.docx"

START_PAGE = 30   # 从第几页开始（1 表示第 1 页）
END_PAGE = 32     # 到第几页结束（先只转 3 页试试，确认效果后再改大）

# ================= 主流程 =================
def main():
    ocr = RapidOCR()
    doc = Document()

    with fitz.open(PDF_PATH) as pdf:
        total = min(END_PAGE, len(pdf))
        print(f"共 {len(pdf)} 页，本次转第 {START_PAGE}~{total} 页\n")
        for i in range(START_PAGE - 1, total):
            page = pdf[i]
            # dpi 越高越清晰、识别越准，但越慢；200 是平衡值
            pix = page.get_pixmap(dpi=200)
            img_bytes = pix.tobytes("png")

            result, _ = ocr(img_bytes)
            page_text = "\n".join([line[1] for line in result]) if result else ""
            doc.add_heading(f"第 {i + 1} 页", level=2)
            doc.add_paragraph(page_text)
            print(f"第 {i + 1} 页完成，识别出 {len(page_text)} 字")

    doc.save(OUT_PATH)
    print(f"\n完成！已保存到：{OUT_PATH}")

if __name__ == "__main__":
    main()
