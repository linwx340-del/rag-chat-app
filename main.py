# -*- coding: utf-8 -*-
"""
RAG 向量检索版（bge embedding + FAISS + DeepSeek）
把之前的 TF-IDF 关键词检索替换成 bge 语义向量检索。

使用前安装依赖：
    pip install sentence-transformers faiss-cpu openai
    （有 NVIDIA 显卡可把 faiss-cpu 换成 faiss-gpu）

使用步骤：
    1. 设置环境变量 DEEPSEEK_API_KEY 为你的真实 key（不要写进代码，防止泄漏）
    2. 把文档 .txt / .pdf / .docx 放进同目录下的 docs 文件夹（脚本会自动读取）
    3. 直接运行，输入问题即可
"""
import os
os.environ["HF_ENDPOINT"] = "https://hf-mirror.com"
import re
import glob
import faiss
from sentence_transformers import SentenceTransformer
from openai import OpenAI

# ================= 配置 =================
DEEPSEEK_API_KEY = os.getenv("DEEPSEEK_API_KEY", "")   # 从环境变量读取，不要写进代码
EMBED_MODEL_NAME = "BAAI/bge-small-zh-v1.5"        # bge 中文向量模型（小版，够用）
CHUNK_SIZE = 400                                   # 每块 400 字
CHUNK_OVERLAP = 50                                 # 相邻块重叠 50 字
TOP_K = 4                                          # 每次检索最相关的 4 块
DOCS_DIR = "./docs"                                # 文档文件夹，.txt / .pdf / .docx 都放这里

# ================= 0. 读取文档 =================
def read_pdf(path):
    """读取 PDF 文件，返回纯文本"""
    try:
        import pdfplumber
    except ImportError:
        raise ImportError("缺少 pdfplumber，请先运行：pip install pdfplumber")
    parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                parts.append(t)
    return "\n".join(parts)


def read_docx(path):
    """读取 Word 文档（.docx），返回纯文本"""
    try:
        import docx
    except ImportError:
        raise ImportError("缺少 python-docx，请先运行：pip install python-docx")
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    # 顺便把表格里的文字也读出来
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def load_docs_folder(folder):
    """读取文件夹里所有 .txt / .pdf / .docx 文件，返回文档内容列表"""
    txt_paths = glob.glob(os.path.join(folder, "*.txt"))
    pdf_paths = glob.glob(os.path.join(folder, "*.pdf"))
    docx_paths = glob.glob(os.path.join(folder, "*.docx"))
    paths = sorted(txt_paths + pdf_paths + docx_paths)

    docs, names = [], []
    for path in paths:
        name = os.path.basename(path)
        try:
            ext = path.lower().rsplit(".", 1)[-1]
            if ext == "pdf":
                text = read_pdf(path)
            elif ext == "docx":
                text = read_docx(path)
            else:
                with open(path, "r", encoding="utf-8") as f:
                    text = f.read()
            text = text.strip()
            if text:
                docs.append(text)
                names.append(name)
            else:
                print(f"  跳过（没提取到文字）：{name}")
        except Exception as e:
            print(f"  读取失败 {name}：{e}")

    print(f"已读取 {len(docs)} 个文档：{names}")
    return docs

# ================= 1. 切分 chunking =================
def chunk_text(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    """按句子把文本切成约 chunk_size 字的小块，相邻块重叠 overlap 字"""
    sentences = re.split(r"(?<=[。！？!?])", text)
    sentences = [s.strip() for s in sentences if s.strip()]

    chunks, cur = [], ""
    for s in sentences:
        if cur and len(cur) + len(s) > chunk_size:
            chunks.append(cur)
            cur = cur[-overlap:] + s   # 新块带上上一块结尾，形成重叠
        else:
            cur += s
    if cur:
        chunks.append(cur)
    return chunks

# ================= 2. 离线入库 =================
def build_index(docs):
    # 加载 bge 向量模型
    embedder = SentenceTransformer(EMBED_MODEL_NAME)
    # 所有文档切块
    all_chunks = []
    for doc in docs:
        all_chunks.extend(chunk_text(doc))
    # 向量化（normalize 之后，内积就等于余弦相似度）
    vectors = embedder.encode(all_chunks, normalize_embeddings=True)
    # 建 FAISS 索引
    dim = vectors.shape[1]
    index = faiss.IndexFlatIP(dim)   # IP = 内积，配合归一化向量做余弦相似度
    index.add(vectors)
    return embedder, index, all_chunks

# ================= 3. 在线查询（检索） =================
def retrieve(embedder, index, chunks, query, top_k=TOP_K):
    # 问题向量化
    q_vec = embedder.encode([query], normalize_embeddings=True)
    # FAISS 检索 top_k 个最相似块
    scores, ids = index.search(q_vec, top_k)
    results = []
    for score, i in zip(scores[0], ids[0]):
        if i >= 0:   # -1 表示没检索到
            results.append((score, chunks[i]))
    return results

# ================= 4. 生成（DeepSeek） =================
def generate(client, query, results):
    context = "\n\n".join(
        f"[资料{idx + 1}] {text}" for idx, (_, text) in enumerate(results)
    )
    prompt = f"""请只根据下面提供的资料回答问题，资料里没有的内容就明确说不知道，不要编造。

资料：
{context}

问题：{query}

回答："""
    resp = client.chat.completions.create(
        model="deepseek-chat",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.3,
    )
    return resp.choices[0].message.content

# ================= 主流程 =================
def main():
    if not DEEPSEEK_API_KEY:
        print("未设置 DeepSeek API Key，请先设置环境变量 DEEPSEEK_API_KEY，例如：")
        print('  PowerShell: $env:DEEPSEEK_API_KEY="sk-你的key"')
        print('  CMD:        set DEEPSEEK_API_KEY=sk-你的key')
        return
    client = OpenAI(api_key=DEEPSEEK_API_KEY, base_url="https://api.deepseek.com")

    # 从 docs 文件夹读文档（txt / pdf 都支持）
    docs = load_docs_folder(DOCS_DIR)
    if not docs:
        print("docs 文件夹里没有文档，请先把 .txt 或 .pdf 放进去再运行！")
        return

    embedder, index, chunks = build_index(docs)
    print(f"已入库 {len(chunks)} 个文本块\n")

    while True:
        query = input("请输入问题（输入 quit 退出）：").strip()
        if query.lower() == "quit":
            break
        if not query:
            continue

        results = retrieve(embedder, index, chunks, query)
        print(f"\n检索到 {len(results)} 块，相似度：", [round(s, 4) for s, _ in results])
        answer = generate(client, query, results)
        print(f"\n回答：{answer}\n" + "=" * 50 + "\n")

if __name__ == "__main__":
    main()